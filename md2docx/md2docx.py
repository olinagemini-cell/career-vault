#!/usr/bin/env python3
"""Convert Markdown resume files to clean Word (.docx) documents.

Target formatting (matching reference resume):
  - Margins: 1.27 cm (0.5 in) all sides
  - Font: Times New Roman throughout (English-only default)
  - Body: 10.5 pt, compact spacing
  - H1 (Name): centered, 22 pt, bold
  - H2 (Section): 16 pt, bold
  - H3 (Company): 14 pt, bold
  - Contact line: centered, 10 pt (first paragraph after H1)

Usage:
    python md2docx.py document.md
    python md2docx.py document.md output.docx
    python md2docx.py *.md
"""

import sys
import glob
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from markdown import markdown
from bs4 import BeautifulSoup, NavigableString


# ── constants ──────────────────────────────────────────────────────────────────
FONT_EN = "Times New Roman"
BODY_SIZE = Pt(10.5)
CONTACT_SIZE = Pt(10)
H1_SIZE = Pt(22)
H2_SIZE = Pt(16)
H3_SIZE = Pt(14)
MARGIN = Cm(1.27)


# ── helpers ──────────────────────────────────────────────────────────────────

def set_fonts(run, en=FONT_EN):
    """Set east-asia + ascii + hAnsi fonts on a run's <w:rPr>."""
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), en)


def paragraph_spacing(para, before=0, after=0, line_spacing=1.08):
    """Set paragraph spacing in points."""
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line_spacing


def spacer(doc, pts=2):
    """Add a visual spacer paragraph (preserved on save)."""
    p = doc.add_paragraph()
    r = p.add_run(" ")
    set_fonts(r)
    r.font.size = Pt(pts)
    paragraph_spacing(p, before=0, after=0, line_spacing=1.0)
    return p


def add_run(para, text, size=BODY_SIZE, bold=False, italic=False,
            color=None, underline=False, en=FONT_EN):
    """Append a fully-styled run to a paragraph. Returns the run."""
    r = para.add_run(text)
    set_fonts(r, en=en)
    r.font.size = size
    r.font.bold = bold
    r.font.italic = italic
    if color:
        r.font.color.rgb = RGBColor(*color)
    r.font.underline = underline
    return r


# ── inline renderer ──────────────────────────────────────────────────────────

def render_inlines(para, element, list_prefix="", en=FONT_EN):
    """Recursively render inline HTML elements into paragraph runs."""
    if list_prefix:
        add_run(para, list_prefix, en=en)

    for child in element.children:
        if isinstance(child, NavigableString):
            text = str(child)
            if text:
                add_run(para, text, en=en)

        elif child.name in ("strong", "b"):
            add_run(para, child.get_text(), bold=True, en=en)

        elif child.name in ("em", "i"):
            add_run(para, child.get_text(), italic=True, en=en)

        elif child.name == "code":
            add_run(para, child.get_text(), size=Pt(9), en="Consolas")

        elif child.name == "del":
            add_run(para, child.get_text(), en=en)

        elif child.name == "a":
            add_run(para, child.get_text(),
                    color=(0x05, 0x63, 0xC1), underline=True, en=en)

        elif child.name == "br":
            para.add_run("\n").font.size = Pt(1)

        else:
            render_inlines(para, child, en=en)


# ── block processor ──────────────────────────────────────────────────────────

def process_blocks(doc, parent, counters, en=FONT_EN):
    """Walk block-level children and add to the document.

    `counters` tracks state:
      - "after_h1": next <p> is the contact line (centered, small).
      - "after_heading": next <p> is first body after a heading (extra space above).
    """
    for child in parent.children:
        if isinstance(child, NavigableString):
            continue
        tag = child.name
        if tag is None:
            continue

        # ── headings ──────────────────────────────────────────────────────
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            level = int(tag[1])
            size = {1: H1_SIZE, 2: H2_SIZE, 3: H3_SIZE,
                    4: Pt(12), 5: Pt(11), 6: Pt(10)}[level]
            before = {1: 0, 2: 12, 3: 12, 4: 10, 5: 8, 6: 8}[level]
            after = {1: 4, 2: 6, 3: 4, 4: 4, 5: 4, 6: 4}[level]

            p = doc.add_paragraph()
            add_run(p, child.get_text().strip(), size=size, bold=True, en=en)
            paragraph_spacing(p, before=before, after=after)

            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                counters["after_h1"] = True
            else:
                counters["after_heading"] = True

        # ── paragraph ─────────────────────────────────────────────────────
        elif tag == "p":
            text = child.get_text().strip()
            if not text:
                spacer(doc)
                counters.pop("after_h1", None)
                counters.pop("after_heading", None)
                continue

            p = doc.add_paragraph()

            if counters.pop("after_h1", False):
                # Contact line
                render_inlines(p, child, en=en)
                for r in p.runs:
                    r.font.size = CONTACT_SIZE
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph_spacing(p, before=0, after=6)

            elif counters.pop("after_heading", False):
                # First body paragraph after a heading
                render_inlines(p, child, en=en)
                paragraph_spacing(p, before=6, after=3)

            else:
                # Continuation body (tight)
                render_inlines(p, child, en=en)
                paragraph_spacing(p, before=3, after=2)

        # ── unordered list ────────────────────────────────────────────────
        elif tag == "ul":
            depth = counters.get("ul_depth", 0) + 1
            counters["ul_depth"] = depth
            for li in child.find_all("li", recursive=False):
                p = doc.add_paragraph()
                render_inlines(p, li, list_prefix="•   ", en=en)
                p.paragraph_format.left_indent = Cm(0.5 * depth)
                paragraph_spacing(p, before=1, after=1)
            counters["ul_depth"] = depth - 1
            counters.pop("after_h1", None)
            counters.pop("after_heading", None)

        # ── ordered list ──────────────────────────────────────────────────
        elif tag == "ol":
            cnt_ref = counters.setdefault("ol_stack", [])
            cnt_ref.append(1)
            for li in [c for c in child.children if c.name == "li"]:
                p = doc.add_paragraph()
                render_inlines(p, li, list_prefix=f"{cnt_ref[-1]}.  ", en=en)
                p.paragraph_format.left_indent = Cm(0.5 * len(cnt_ref))
                paragraph_spacing(p, before=1, after=1)
                cnt_ref[-1] += 1
            cnt_ref.pop()
            counters.pop("after_h1", None)
            counters.pop("after_heading", None)

        # ── blockquote ────────────────────────────────────────────────────
        elif tag == "blockquote":
            before_idx = len(doc.paragraphs)
            process_blocks(doc, child, counters, en=en)
            for par in doc.paragraphs[before_idx:]:
                par.paragraph_format.left_indent = Cm(0.75)

        # ── code block ────────────────────────────────────────────────────
        elif tag == "pre":
            code = child.get_text()
            if not code:
                continue
            for line in code.split("\n"):
                p = doc.add_paragraph()
                add_run(p, line or " ", size=Pt(9), en="Consolas")
                paragraph_spacing(p, before=0, after=0, line_spacing=1.0)
                p.paragraph_format.left_indent = Cm(0.5)

        # ── horizontal rule ───────────────────────────────────────────────
        elif tag == "hr":
            p = doc.add_paragraph()
            add_run(p, "─" * 60, size=Pt(8), color=(0xCC, 0xCC, 0xCC))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph_spacing(p, before=8, after=8)

        # ── table ─────────────────────────────────────────────────────────
        elif tag == "table":
            rows = child.find_all("tr")
            if not rows:
                continue
            max_cells = max(len(r.find_all(["td", "th"])) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=max_cells)
            tbl.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, cell in enumerate(row.find_all(["td", "th"])):
                    tc = tbl.cell(i, j)
                    tc.text = cell.get_text().strip()
                    for tp in tc.paragraphs:
                        for tr in tp.runs:
                            set_fonts(tr, en=en)
            spacer(doc)

        # ── recurse ───────────────────────────────────────────────────────
        else:
            process_blocks(doc, child, counters, en=en)


# ── main converter ───────────────────────────────────────────────────────────

def md_to_docx(md_path, docx_path=None):
    """Convert `md_path` to a .docx file. Returns the output path."""
    md_path = Path(md_path)
    if not md_path.is_file():
        raise FileNotFoundError(f"File not found: {md_path}")

    if docx_path is None:
        docx_path = md_path.with_suffix(".docx")

    # Markdown → HTML
    md_text = md_path.read_text(encoding="utf-8")
    html = markdown(md_text, extensions=["tables", "fenced_code"])
    soup = BeautifulSoup(html, "lxml")
    body_elem = soup.find("body") or soup

    # Build Word document
    doc = Document()

    # Normal style: Times New Roman, 10.5 pt
    style = doc.styles["Normal"]
    style.font.name = FONT_EN
    style.font.size = BODY_SIZE
    srPr = style.element.get_or_add_rPr()
    srf = srPr.find(qn("w:rFonts"))
    if srf is None:
        srf = srPr.makeelement(qn("w:rFonts"), {})
        srPr.insert(0, srf)
    srf.set(qn("w:ascii"), FONT_EN)
    srf.set(qn("w:hAnsi"), FONT_EN)
    srf.set(qn("w:eastAsia"), FONT_EN)

    # Page margins: 1.27 cm all sides
    for sec in doc.sections:
        sec.top_margin = MARGIN
        sec.bottom_margin = MARGIN
        sec.left_margin = MARGIN
        sec.right_margin = MARGIN

    process_blocks(doc, body_elem, {}, en=FONT_EN)

    doc.save(str(docx_path))
    return docx_path


# ── cli ──────────────────────────────────────────────────────────────────────

def main():
    if len(sys.argv) < 2:
        print("Usage: python md2docx.py <file.md> [output.docx]")
        print("       python md2docx.py *.md")
        sys.exit(1)

    args = sys.argv[1:]
    md_files = []
    output_path = None
    for a in args:
        if a.endswith(".md"):
            md_files.extend(glob.glob(a) or [a])
        elif a.endswith(".docx") and output_path is None:
            output_path = a

    for f in md_files:
        out = md_to_docx(f, output_path if len(md_files) == 1 else None)
        print(f"  [OK]  {f}  ->  {out}")

    if not md_files:
        print("No .md files found.")


if __name__ == "__main__":
    main()
